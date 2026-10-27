import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import openpyxl
import datetime
import os
import glob
import json
import zipfile
import io
import time

st.set_page_config(page_title="発注検討表 ＆ 弥生受注管理システム", layout="wide")

st.markdown("""
<style>
    header[data-testid="stHeader"] {
        height: 1.8rem !important;
        background: transparent !important;
    }
    .block-container {
        padding-top: 1.5rem !important;
        padding-bottom: 1rem !important;
        padding-left: 0.8rem !important;
        padding-right: 0.8rem !important;
        max-width: 100% !important;
    }
    .app-header {
        font-family: "Meiryo", sans-serif;
        font-size: 20px;
        font-weight: bold;
        color: #1f2937;
        margin-bottom: 6px;
        padding-bottom: 4px;
        border-bottom: 3px solid #2563eb;
    }
</style>
""", unsafe_allow_html=True)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__)) if "__file__" in locals() else "."

SEARCH_DIRS = [
    SCRIPT_DIR,
    ".",
    r"G:\マイドライブ\発注検討",
    r"G:\マイドライブ\売上ダッシュボード\済",
    r"G:\マイドライブ\売上ダッシュボード",
    "/content/drive/MyDrive/発注検討",
    "/content/drive/MyDrive/売上ダッシュボード/済"
]

def find_target_file(keywords, ext=".xlsx"):
    for d in SEARCH_DIRS:
        if os.path.exists(d):
            try:
                for fn in os.listdir(d):
                    if fn.lower().endswith(ext) and not fn.startswith("~$"):
                        if any(k in fn for k in keywords):
                            return os.path.join(d, fn)
            except Exception:
                pass
    return None

MAIN_EXCEL = find_target_file(["発注検討表"])
YAYOI_EXCEL = find_target_file(["弥生確保リスト", "確保リスト"])
YAYOI_ITEM_EXCEL = find_target_file(["弥生商品番号", "商品番号"])
YAYOI_COMP_EXCEL = find_target_file(["弥生会社名", "会社名", "得意先"])

SAVE_CSV = os.path.join(SCRIPT_DIR, "meeting_decisions_v8.csv")
MANUAL_YAYOI_CSV = os.path.join(SCRIPT_DIR, "yayoi_manual_additions.csv")

STAFF_OPTIONS_CLEAN = ["", "小泉", "鈴木", "松平", "原", "西山", "田中", "水原", "川村", "穴山"]

MARUFUJI_TARGET_NAMES = {
    "株式会社 丸藤",
    "株式会社 丸藤 九州営業グループ",
    "株式会社 丸藤 東日本",
    "株式会社 丸藤 東日本",
    "株式会社丸藤",
    "株式会社丸藤東日本"
}

def is_na_val(val):
    if val is None or pd.isna(val):
        return True
    s = str(val).strip().lower()
    return s in ("#n/a", "#ref!", "#div/0!", "#value!", "#name?", "none", "nan", "<na>", "")

def clean_str(val):
    if is_na_val(val):
        return ""
    s = str(val).strip()
    if s.endswith(".0") and s[:-2].replace("-", "").isdigit():
        return s[:-2]
    return s

def normalize_code(val):
    s = clean_str(val).upper()
    return s.replace(" ", "").replace(" ", "").replace("－", "-").replace("ー", "-").replace("―", "-")

def to_num(val):
    if is_na_val(val):
        return 0.0
    try:
        s = str(val).replace(",", "").replace("，", "").replace("%", "").replace("円", "").replace("(", "-").replace(")", "").strip()
        return float(s)
    except Exception:
        return 0.0

@st.cache_data
def load_company_master(file_source):
    loaded = []
    seen = set()
    sources = [file_source, YAYOI_COMP_EXCEL, "弥生会社名.xlsx", r"G:\マイドライブ\発注検討\弥生会社名.xlsx"]
    for src in sources:
        if src and os.path.exists(src):
            try:
                df_c = pd.read_excel(src, sheet_name=0, header=None)
                for idx in range(len(df_c)):
                    r = df_c.iloc[idx].values
                    if len(r) < 3:
                        continue
                    c_code = clean_str(r[1])
                    c_name = clean_str(r[2])
                    if not c_name or c_name in ("名称", "得意先名", "会社名", "得意先リスト", "コード順", "コード"):
                        continue
                    if c_name not in seen:
                        seen.add(c_name)
                        loaded.append({"code": c_code, "name": c_name})
                if len(loaded) > 50:
                    break
            except Exception:
                continue
    return loaded

@st.cache_data
def load_item_master(file_source):
    loaded = []
    seen = set()
    sources = [file_source, YAYOI_ITEM_EXCEL, "弥生商品番号.xlsx", r"G:\マイドライブ\発注検討\弥生商品番号.xlsx"]
    for src in sources:
        if src and os.path.exists(src):
            try:
                df_i = pd.read_excel(src, sheet_name=0, header=None)
                for idx in range(len(df_i)):
                    r = df_i.iloc[idx].values
                    if len(r) < 3:
                        continue
                    i_code = clean_str(r[1])
                    i_name = clean_str(r[2])
                    i_wh = clean_str(r[20]) if len(r) > 20 else ""
                    if not i_code or i_code in ("コード", "商品コード", "品番", "商品リスト", "コード順", "名称"):
                        continue
                    if i_code not in seen:
                        seen.add(i_code)
                        loaded.append({"code": i_code, "name": i_name, "wh": i_wh if i_wh in ("青山", "名古屋", "共通") else ""})
                if len(loaded) > 50:
                    break
            except Exception:
                continue
    return loaded

@st.cache_data
def load_excel_master(file_source):
    if file_source is None:
        return pd.DataFrame(), pd.DataFrame(), "前回決定事項"
    target_df = None
    all_sheets = {}
    try:
        all_sheets = pd.read_excel(file_source, sheet_name=None, header=None)
        target_df = list(all_sheets.values())[0]
    except Exception as e:
        try:
            wb = openpyxl.load_workbook(file_source, data_only=True, keep_vba=False)
            ws = wb[wb.sheetnames[0]]
            data = [[cell.value for cell in row] for row in ws.iter_rows()]
            target_df = pd.DataFrame(data)
            all_sheets = {wb.sheetnames[0]: target_df}
        except Exception as ex:
            st.error(f"発注検討表の読み込みエラー: {ex}")
            return pd.DataFrame(), pd.DataFrame(), "前回決定事項"

    header_idx = 0
    for idx in range(min(15, len(target_df))):
        row_vals = [str(x) for x in target_df.iloc[idx].values]
        if any("品番" in x for x in row_vals):
            header_idx = idx
            break

    h_row = [str(x).replace("\n", "").replace(" ", "").replace(" ", "") for x in target_df.iloc[header_idx].values]
    col_map = {}
    for i, c in enumerate(h_row):
        if "商品センター" in c and "center" not in col_map: col_map["center"] = i
        elif "品番" in c and "code" not in col_map: col_map["code"] = i
        elif "商品区分" in c and "name" not in col_map: col_map["name"] = i
        elif "73期" in c and "k73" not in col_map: col_map["k73"] = i
        elif "74期" in c and "k74" not in col_map: col_map["k74"] = i
        elif "75期" in c and "k75" not in col_map: col_map["k75"] = i
        elif "76期" in c and "k76" not in col_map: col_map["k76"] = i
        elif "77期" in c and "k77" not in col_map: col_map["k77"] = i
        elif ("2課発注点" in c or "２課発注点" in c) and "point2" not in col_map: col_map["point2"] = i
        elif "発注点" in c and "point1" not in col_map: col_map["point1"] = i
        elif "在庫数" in c and "roji_stock" not in col_map: col_map["roji_stock"] = i
        elif "入荷予定日" in c and "ndate" not in col_map: col_map["ndate"] = i
        elif "入荷予定" in c and "roji_in" not in col_map: col_map["roji_in"] = i
        elif "ロジザード" in c and ("確保" in c or "在庫確保" in c) and "roji_sec" not in col_map: col_map["roji_sec"] = i
        elif ("フリー在庫" in c or "有効数量" in c) and "free" not in col_map: col_map["free"] = i
        elif "回転率" in c and "rate" not in col_map: col_map["rate"] = i
        elif ("リンベル" in c or "ﾘﾝﾍﾞﾙ" in c) and "rinbell" not in col_map: col_map["rinbell"] = i
        elif "前回決定事項" in c and "prev_dec" not in col_map: col_map["prev_dec"] = i
        elif "決定事項" in c and "curr_dec" not in col_map: col_map["curr_dec"] = i

    c_center = col_map.get("center", 1)
    c_code = col_map.get("code", 2)
    c_name = col_map.get("name", 3)
    c_k73 = col_map.get("k73", 6)
    c_k74 = col_map.get("k74", 7)
    c_k75 = col_map.get("k75", 9)
    c_k76 = col_map.get("k76", 10)
    c_point1 = col_map.get("point1", 14)
    c_point2 = col_map.get("point2", -1)
    c_roji_stock = col_map.get("roji_stock", 15)
    c_roji_in = col_map.get("roji_in", 16)
    c_ndate = col_map.get("ndate", 18)
    c_roji_sec = col_map.get("roji_sec", 19)
    c_free = col_map.get("free", 23)
    c_rate = col_map.get("rate", 25)
    c_rin = col_map.get("rinbell", 26)
    c_prev = col_map.get("prev_dec", 27)
    c_curr = col_map.get("curr_dec", 28)

    raw_h = target_df.iloc[header_idx].values
    prev_col_title = clean_str(raw_h[c_prev]) if c_prev < len(raw_h) and clean_str(raw_h[c_prev]) else "前回決定事項"

    items = []
    for idx in range(header_idx + 1, len(target_df)):
        r = target_df.iloc[idx].values
        if len(r) <= c_code:
            continue
        hinban = clean_str(r[c_code])
        if not hinban or hinban == "品番":
            continue

        p1_val = int(round(to_num(r[c_point1]), 0)) if c_point1 < len(r) and not is_na_val(r[c_point1]) else 0
        p2_val = int(round(to_num(r[c_point2]), 0)) if 0 <= c_point2 < len(r) and not is_na_val(r[c_point2]) else 0
        free_val = int(to_num(r[c_free])) if c_free < len(r) and not is_na_val(r[c_free]) else 0
        flag_val = free_val - (p1_val + p2_val)

        sec_roji = int(to_num(r[c_roji_sec])) if c_roji_sec < len(r) else 0

        rate_raw = to_num(r[c_rate]) if c_rate < len(r) else 0.0
        rate_val = int(round(rate_raw * 100, 0)) if abs(rate_raw) <= 5.0 else int(round(rate_raw, 0))

        ndate_val = clean_str(r[c_ndate]) if c_ndate < len(r) else ""
        if "00:00:00" in ndate_val:
            ndate_val = ndate_val.split(" ")[0]

        items.append({
            "No": int(to_num(r[0])) if not is_na_val(r[0]) else idx,
            "商品センター": clean_str(r[c_center]) if c_center < len(r) else "",
            "品番": hinban,
            "品番_正規化": normalize_code(hinban),
            "商品区分": clean_str(r[c_name]) if c_name < len(r) else "",
            "73期": int(to_num(r[c_k73])) if c_k73 < len(r) else 0,
            "74期": int(to_num(r[c_k74])) if c_k74 < len(r) else 0,
            "75期": int(to_num(r[c_k75])) if c_k75 < len(r) else 0,
            "76期": int(to_num(r[c_k76])) if c_k76 < len(r) else 0,
            "1課発注点": p1_val,
            "2課発注点": p2_val,
            "在庫数_ロジザード": int(to_num(r[c_roji_stock])) if c_roji_stock < len(r) else 0,
            "入荷予定数_ロジザード": int(to_num(r[c_roji_in])) if c_roji_in < len(r) else 0,
            "入荷予定日_ロジザード": ndate_val,
            "確保数_ロジザード": sec_roji,
            "在庫確保_弥生": 0,
            "在庫確保_丸藤": 0,
            "確保数のズレ": sec_roji,
            "入荷後のフリー在庫": free_val,
            "発注検討フラグ": flag_val,
            "在庫回転率": rate_val,
            "リンベル確保数": clean_str(r[c_rin]) if c_rin < len(r) else "",
            "前回決定事項": clean_str(r[c_prev]) if c_prev < len(r) else "",
            "決定事項": clean_str(r[c_curr]) if c_curr < len(r) else "",
            "発注日": "",
            "NA行": False
        })

    df_main = pd.DataFrame(items)

    hist_list = []
    for s_name, df_s in all_sheets.items():
        s_str = str(s_name).strip()
        if s_str.isdigit() and len(s_str) == 5 and not df_s.empty:
            h_row_vals = [clean_str(x) for x in df_s.iloc[0].values]
            h_col_idx = next((i for i, c in enumerate(h_row_vals) if c == "品番"), 2)
            k_cols = [i for i, c in enumerate(h_row_vals) if c and "決定事項" in c and "前回" not in c]
            s_clean = s_str.lstrip("0")
            rest = s_clean[1:] if s_clean.startswith("8") else s_clean
            dt_str = f"2026-{rest[:2]}-{rest[2:]}" if len(rest) == 4 else (f"2026-0{rest[0]}-{rest[1:]}" if len(rest) == 3 else s_str)
            for _, r_d in df_s.iloc[1:].iterrows():
                r_vals = r_d.values
                if len(r_vals) <= h_col_idx:
                    continue
                c_val = clean_str(r_vals[h_col_idx])
                if not c_val:
                    continue
                for kc in k_cols:
                    if kc < len(r_vals) and clean_str(r_vals[kc]) not in ("", "0"):
                        hist_list.append({"会議日": dt_str, "シート": s_str, "品番": c_val, "決定事項": clean_str(r_vals[kc])})

    df_hist = pd.DataFrame(hist_list).drop_duplicates(subset=["会議日", "品番", "決定事項"]) if hist_list else pd.DataFrame(columns=["会議日", "シート", "品番", "決定事項"])
    return df_main, df_hist, prev_col_title

@st.cache_data
def load_yayoi_data(file_source):
    df = pd.DataFrame()
    sources = [file_source, YAYOI_EXCEL, "弥生確保リスト.xlsx", r"G:\マイドライブ\発注検討\弥生確保リスト.xlsx"]
    for src in sources:
        if src and os.path.exists(src):
            try:
                df_raw = pd.read_excel(src, sheet_name=0, header=None)
                h_idx = 4
                for i in range(min(15, len(df_raw))):
                    row_v = [clean_str(x) for x in df_raw.iloc[i].values]
                    if "商品コード" in row_v or "得意先名" in row_v:
                        h_idx = i
                        break
                
                df = pd.read_excel(src, sheet_name=0, header=h_idx)
                if not df.empty and ("商品コード" in df.columns or len(df.columns) > 21):
                    break
            except Exception:
                continue

    if not df.empty:
        col_c_code = "商品コード" if "商品コード" in df.columns else df.columns[13]
        col_c_qty = "数量" if "数量" in df.columns else df.columns[21]
        col_c_cust = "得意先名" if "得意先名" in df.columns else df.columns[5]
        col_c_ord = "受注番号" if "受注番号" in df.columns else df.columns[2]
        col_c_odate = "受注日" if "受注日" in df.columns else df.columns[1]
        col_c_ddate = "納入期日" if "納入期日" in df.columns else (df.columns[32] if len(df.columns) > 32 else df.columns[1])
        col_c_wh = "倉庫名" if "倉庫名" in df.columns else (df.columns[19] if len(df.columns) > 19 else "")
        col_c_staff = "担当者名" if "担当者名" in df.columns else (df.columns[10] if len(df.columns) > 10 else "")
        col_c_note = "備考" if "備考" in df.columns else (df.columns[30] if len(df.columns) > 30 else "")

        df = df[df[col_c_code].notna()].copy()
        df["商品コード"] = df[col_c_code].apply(clean_str)
        df = df[df["商品コード"] != ""]
        df["得意先名"] = df[col_c_cust].apply(lambda x: clean_str(x) if clean_str(x) else "不明")
        df["数量"] = pd.to_numeric(df[col_c_qty], errors="coerce").fillna(0).astype(int)
        df["受注残数"] = df["数量"]
        df["受注番号"] = df[col_c_ord].apply(clean_str)
        df["受注日_dt"] = pd.to_datetime(df[col_c_odate], errors="coerce")
        df["納入期日_dt"] = pd.to_datetime(df[col_c_ddate], errors="coerce")
        df["受注日"] = df["受注日_dt"].dt.strftime("%Y-%m-%d").fillna("")
        df["納入期日"] = df["納入期日_dt"].dt.strftime("%Y-%m-%d").fillna("")
        df["倉庫名"] = df[col_c_wh].apply(clean_str) if col_c_wh in df.columns else ""
        df["担当者名"] = df[col_c_staff].apply(clean_str) if col_c_staff in df.columns else ""
        df["備考"] = df[col_c_note].apply(clean_str) if col_c_note in df.columns else ""

        empty_mask = df["受注番号"] == ""
        if empty_mask.any():
            df.loc[empty_mask, "受注番号"] = df[empty_mask].apply(
                lambda r: f"Y-{r['受注日'].replace('-','')}-{abs(hash(r['得意先名']))%10000:04d}", axis=1
            )
    else:
        df = pd.DataFrame(columns=["受注番号", "得意先名", "商品コード", "数量", "受注残数", "受注日", "納入期日", "備考", "倉庫名", "担当者名", "受注日_dt", "納入期日_dt"])

    if os.path.exists(MANUAL_YAYOI_CSV):
        try:
            df_man = pd.read_csv(MANUAL_YAYOI_CSV).fillna("")
            if not df_man.empty:
                df_man["数量"] = pd.to_numeric(df_man["数量"], errors="coerce").fillna(0).astype(int)
                df_man["受注残数"] = df_man["数量"]
                df_man["受注日_dt"] = pd.to_datetime(df_man["受注日"], errors="coerce")
                df_man["納入期日_dt"] = pd.to_datetime(df_man["納入期日"], errors="coerce")
                edited_onos = set(df_man["受注番号"].astype(str).unique())
                df = df[~df["受注番号"].astype(str).isin(edited_onos)]
                df = pd.concat([df, df_man], ignore_index=True)
        except Exception:
            pass

    if not df.empty:
        df = df.sort_values(by=["受注日_dt", "納入期日_dt"], ascending=[False, False], na_position="last")
    return df

@st.cache_data
def load_sales_csvs_universal(uploaded_files, valid_norm_codes_tuple):
    valid_codes = set(valid_norm_codes_tuple)
    all_records = []

    processed_files = []
    if uploaded_files:
        for f in uploaded_files:
            fname = getattr(f, "name", "uploaded.csv")
            if fname.lower().endswith(".zip"):
                try:
                    f.seek(0)
                    with zipfile.ZipFile(io.BytesIO(f.read())) as z:
                        for zname in z.namelist():
                            if zname.lower().endswith(".csv") and not zname.startswith("__MACOSX") and not zname.startswith("._"):
                                with z.open(zname) as zf:
                                    processed_files.append((zname, io.BytesIO(zf.read())))
                except Exception as ze:
                    st.error(f"ZIP解凍エラー ({fname}): {ze}")
            else:
                f.seek(0)
                processed_files.append((fname, f))

    local_csv_paths = []
    for d in SEARCH_DIRS:
        if os.path.exists(d):
            local_csv_paths.extend(glob.glob(os.path.join(d, "yayoi_*.csv")))
            local_csv_paths.extend(glob.glob(os.path.join(d, "*yayoi*.csv")))
            local_csv_paths.extend(glob.glob(os.path.join(d, "*.csv")))
    local_csv_paths = sorted(list(set(local_csv_paths)))

    if not processed_files and not local_csv_paths:
        return pd.Series(dtype=int), 0, pd.DataFrame(), {}

    sources = processed_files if processed_files else [(os.path.basename(p), p) for p in local_csv_paths]
    loaded_file_count = 0
    period_buckets = {"今月（最新CSV）": {}}

    for fname, f in sources:
        df_s = None
        for enc in ["cp932", "shift_jis", "utf-8-sig", "utf-8"]:
            try:
                if isinstance(f, str):
                    df_s = pd.read_csv(f, encoding=enc, header=None, dtype=str)
                else:
                    f.seek(0)
                    df_s = pd.read_csv(f, encoding=enc, header=None, dtype=str)
                if df_s is not None and not df_s.empty:
                    break
            except Exception:
                continue

        if df_s is None or df_s.empty:
            continue

        loaded_file_count += 1
        code_col = 15
        qty_col = 23
        start_row = 0
        for r_i in range(min(5, len(df_s))):
            row_str = [str(x).strip() for x in df_s.iloc[r_i].values]
            if "商品コード" in row_str:
                start_row = r_i + 1
                code_col = row_str.index("商品コード")
                if "数量" in row_str:
                    qty_col = row_str.index("数量")
                break

        file_records = []
        for r_i in range(start_row, len(df_s)):
            row_vals = df_s.iloc[r_i].values
            if len(row_vals) <= max(code_col, qty_col):
                continue

            c_val = normalize_code(row_vals[code_col])
            if c_val in valid_codes:
                q_val = to_num(row_vals[qty_col])
                if q_val > 0:
                    d_val = clean_str(row_vals[3]) if len(row_vals) > 3 else ""
                    all_records.append({
                        "file": fname,
                        "date": d_val,
                        "code_norm": c_val,
                        "qty": q_val
                    })
                    file_records.append({"code_norm": c_val, "qty": q_val})

        if file_records:
            df_fr = pd.DataFrame(file_records)
            f_summary = df_fr.groupby("code_norm")["qty"].sum().round().astype(int).to_dict()
            period_buckets[fname] = f_summary

    if not all_records:
        return pd.Series(dtype=int), 0, pd.DataFrame(), {}

    df_rec = pd.DataFrame(all_records)
    summary_series = df_rec.groupby("code_norm")["qty"].sum().round().astype(int)
    period_buckets["全期間合算（合計）"] = summary_series.to_dict()

    return summary_series, loaded_file_count, df_rec, period_buckets

# ---------------------------------------------------------
# ヘッダー ＆ アップローダー（session_state保持でリセット防止）
# ---------------------------------------------------------
if "persisted_sales_files" not in st.session_state:
    st.session_state["persisted_sales_files"] = None

head_col1, head_col2 = st.columns([3.2, 2.0])
with head_col1:
    st.markdown('<div class="app-header">📊 発注検討表 ＆ 弥生受注管理システム（完全連動版）</div>', unsafe_allow_html=True)
with head_col2:
    with st.expander("📂 売上CSV・ZIP ＆ 最新Excelの読込・更新", expanded=False):
        up_sales_raw = st.file_uploader("🔥 売上CSV または ZIPファイル（複数選択・ZIP一括ドロップ対応）", type=["csv", "zip"], accept_multiple_files=True, key="up_s_raw")
        if up_sales_raw:
            st.session_state["persisted_sales_files"] = up_sales_raw
        
        up_sales = st.session_state["persisted_sales_files"]

        u_c1, u_c2 = st.columns(2)
        up_main = u_c1.file_uploader("📋 発注検討表.xlsx", type=["xlsx"], key="up_m")
        up_yayoi = u_c2.file_uploader("📑 弥生確保リスト.xlsx", type=["xlsx"], key="up_y")
        u_c3, u_c4 = st.columns(2)
        up_comp = u_c3.file_uploader("🏢 弥生会社名.xlsx", type=["xlsx"], key="up_comp")
        up_item = u_c4.file_uploader("📦 弥生商品番号.xlsx", type=["xlsx"], key="up_item")

        col_btn1, col_btn2 = st.columns(2)
        if col_btn1.button("🔄 最新状態に再読み込み", use_container_width=True):
            st.cache_data.clear()
            st.rerun()
        if col_btn2.button("🗑️ 読込CSVをクリア", use_container_width=True):
            st.session_state["persisted_sales_files"] = None
            st.cache_data.clear()
            st.rerun()

main_src = up_main if up_main is not None else MAIN_EXCEL
yayoi_src = up_yayoi if up_yayoi is not None else YAYOI_EXCEL
comp_src = up_comp if up_comp is not None else YAYOI_COMP_EXCEL
item_src = up_item if up_item is not None else YAYOI_ITEM_EXCEL

df_main, df_hist, prev_col_title = load_excel_master(main_src)
df_yayoi = load_yayoi_data(yayoi_src)
master_comps = load_company_master(comp_src)
master_items = load_item_master(item_src)

if df_main.empty:
    st.error("⚠️ 『発注検討表.xlsx』が見つかりません。右上のエクスパンダーから『発注検討表.xlsx』をアップロードしてください。")
    st.stop()

marufuji_dict = {}
yayoi_dict = {}
if not df_yayoi.empty:
    df_yayoi["品番_正規化"] = df_yayoi["商品コード"].apply(normalize_code)
    yayoi_real_totals = df_yayoi.groupby("品番_正規化")["数量"].sum().to_dict()
    
    def is_marufuji_name(name):
        n_clean = str(name).strip().replace(" ", " ")
        return any(target.replace(" ", " ") in n_clean for target in MARUFUJI_TARGET_NAMES)
    
    df_maru_only = df_yayoi[df_yayoi["得意先名"].apply(is_marufuji_name)]
    marufuji_real_totals = df_maru_only.groupby("品番_正規化")["数量"].sum().to_dict()

    df_main["在庫確保_弥生"] = df_main["品番_正規化"].map(yayoi_real_totals).fillna(0).astype(int)
    df_main["在庫確保_丸藤"] = df_main["品番_正規化"].map(marufuji_real_totals).fillna(0).astype(int)
    df_main["確保数のズレ"] = df_main["確保数_ロジザード"] - df_main["在庫確保_弥生"] - df_main["在庫確保_丸藤"]

    for code_n, grp in df_yayoi.groupby("品番_正規化", sort=False):
        grp_sorted = grp.sort_values(by=["受注日_dt", "納入期日_dt"], ascending=[False, False], na_position="last")
        yayoi_dict[code_n] = grp_sorted[["受注番号", "得意先名", "数量", "受注日", "納入期日", "倉庫名", "備考", "担当者名"]].to_dict(orient="records")
        maru_grp = grp_sorted[grp_sorted["得意先名"].apply(is_marufuji_name)]
        if not maru_grp.empty:
            marufuji_dict[code_n] = maru_grp[["受注番号", "得意先名", "数量", "受注日", "納入期日", "倉庫名", "備考", "担当者名"]].to_dict(orient="records")

item_wh_map = {}
item_name_map = {}
for it in master_items:
    c_norm = normalize_code(it["code"])
    raw_wh = clean_str(it.get("wh", ""))
    item_wh_map[c_norm] = raw_wh
    item_wh_map[it["code"]] = raw_wh
    item_name_map[c_norm] = it["name"]
    item_name_map[it["code"]] = it["name"]

seen_i_codes = {normalize_code(it["code"]) for it in master_items}
for _, r in df_main.iterrows():
    ic = clean_str(r["品番"])
    ic_n = normalize_code(ic)
    if ic and ic_n not in seen_i_codes:
        seen_i_codes.add(ic_n)
        wh_def = clean_str(r["商品センター"]) if clean_str(r["商品センター"]) in ("青山", "名古屋") else ""
        master_items.append({"code": ic, "name": clean_str(r["商品区分"]), "wh": wh_def})
        item_wh_map[ic] = wh_def
        item_wh_map[ic_n] = wh_def
        item_name_map[ic] = clean_str(r["商品区分"])

all_valid_norm_codes = tuple(set(df_main["品番_正規化"].unique()) | set(item_name_map.keys()))
sales_series, sales_file_cnt, df_sales_detail, period_buckets = load_sales_csvs_universal(up_sales, all_valid_norm_codes)

# 過去期・期間の切り替えセレクター
period_options = list(period_buckets.keys()) if period_buckets else ["全期間合算（合計）"]
sel_period = st.selectbox("📅 表示する売上期間・CSVファイルの切り替え（過去の期や月別も選択可）", period_options, index=0)

active_sales_dict = period_buckets.get(sel_period, {}) if period_buckets else {}
if active_sales_dict:
    df_main["今月売れ数"] = df_main["品番_正規化"].map(active_sales_dict).fillna(0).astype(int)
else:
    df_main["今月売れ数"] = 0

m_df = df_main.copy()
hist_dict = {} if 'df_hist' in locals() and df_hist is not None and not df_hist.empty else {}

total_sales_qty = int(m_df["今月売れ数"].sum())
asc001_qty = int(active_sales_dict.get("ASC-001", 0)) if active_sales_dict else 0
st.caption(
    f"✅ **最新データ完全連動中** ｜ 発注検討表: **{len(m_df):,}** 商品 ｜ "
    f"得意先マスタ: **{len(master_comps):,}** 社 ｜ "
    f"🔥 選択期間の売れ数合計: **{total_sales_qty:,}** 個 （読込ファイル: {sales_file_cnt} 件 ／ ASC-001: **{asc001_qty:,}** 個）"
)

tab1, tab2, tab3, tab4 = st.tabs([
    f"📋 1. 発注検討表（発注点リアルタイム連動版）",
    "🏢 2. 会社（得意先）別・受注番号別 確保リスト（特大ポップアップ）",
    "💻 3. 弥生受注伝票（コード・商品名・倉庫名 分離版 ／ 最大100行）",
    "🕒 4. 入荷予定・会議履歴 ＆ 売上CSV検証"
])

with tab1:
    f_col1, f_col2, f_col3 = st.columns([2.6, 1.3, 3.0])
    filter_mode = f_col1.selectbox(
        "表示絞り込み",
        [
            "すべて表示（Excelと同じ並び）",
            "🚨 要発注（発注検討フラグがマイナスのみ）",
            "🔥 【ヒット商品】在庫があるのに今月売れている",
            "🔥 今月売れている商品すべて（売れ数1以上）",
            "確保数のズレがある商品",
            "入荷予定がある商品",
            "弥生確保がある商品"
        ],
        index=0
    )
    center_opt = ["すべて"] + sorted([x for x in m_df["商品センター"].unique() if x])
    sel_center = f_col2.selectbox("商品センター", center_opt)
    search_kw = f_col3.text_input("🔍 品番・商品名・会社名・決定事項で検索", placeholder="例: ASC-001 / AHL-2203 / ライト / 丸藤")

    v_df = m_df.copy()
    if filter_mode == "🚨 要発注（発注検討フラグがマイナスのみ）":
        v_df = v_df[(v_df["発注検討フラグ"] < 0) & (~v_df["NA行"])]
    elif filter_mode == "🔥 【ヒット商品】在庫があるのに今月売れている":
        v_df = v_df[(v_df["入荷後のフリー在庫"] >= 1) & (v_df["今月売れ数"] >= 1) & (~v_df["NA行"])]
    elif filter_mode == "🔥 今月売れている商品すべて（売れ数1以上）":
        v_df = v_df[v_df["今月売れ数"] >= 1].sort_values(by="今月売れ数", ascending=False)
    elif filter_mode == "確保数のズレがある商品":
        v_df = v_df[v_df["確保数のズレ"] != 0]
    elif filter_mode == "入荷予定がある商品":
        v_df = v_df[v_df["入荷予定数_ロジザード"] > 0]
    elif filter_mode == "弥生確保がある商品":
        v_df = v_df[(v_df["在庫確保_弥生"] > 0) | (v_df["在庫確保_丸藤"] > 0)]

    if sel_center != "すべて":
        v_df = v_df[v_df["商品センター"] == sel_center]
    if search_kw:
        v_df = v_df[
            v_df["品番"].str.contains(search_kw, case=False, na=False) |
            v_df["商品区分"].str.contains(search_kw, case=False, na=False) |
            v_df["前回決定事項"].str.contains(search_kw, case=False, na=False) |
            v_df["決定事項"].str.contains(search_kw, case=False, na=False)
        ]

    records_json = json.dumps(v_df.to_dict(orient="records"), ensure_ascii=False)
    yayoi_json = json.dumps(yayoi_dict, ensure_ascii=False)
    marufuji_json = json.dumps(marufuji_dict, ensure_ascii=False)
    hist_json = json.dumps(hist_dict, ensure_ascii=False)
    prev_title_html = prev_col_title.replace("(", "<br>(")

    html_tab1 = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <style>
        * {{ box-sizing: border-box; font-family: "Meiryo", "Hiragino Kaku Gothic ProN", sans-serif; }}
        body {{ margin: 0; padding: 0; background: #fff; }}
        .toolbar {{
            display: flex; align-items: center; justify-content: space-between;
            background: #f8fafc; padding: 6px 12px; border: 1px solid #cbd5e1;
            border-bottom: none; border-radius: 6px 6px 0 0;
        }}
        .toolbar-info {{ font-size: 13px; color: #1e293b; font-weight: bold; }}
        .toolbar-hint {{ font-size: 12px; color: #0369a1; background: #e0f2fe; padding: 4px 10px; border-radius: 4px; font-weight: bold; }}
        .btn-download {{
            padding: 6px 14px; font-size: 12px; font-weight: bold; border-radius: 4px;
            cursor: pointer; border: none; background: #2563eb; color: #fff;
        }}
        .btn-download:hover {{ background: #1d4ed8; }}
        .table-wrap {{ width: 100%; height: 620px; overflow: auto; border: 1px solid #94a3b8; position: relative; }}
        table {{ border-collapse: separate; border-spacing: 0; width: max-content; min-width: 100%; }}
        th, td {{ border-right: 1px solid #94a3b8; border-bottom: 1px solid #94a3b8; padding: 3px 6px; font-size: 12px; height: 30px; }}
        thead th {{
            position: sticky; top: 0; background-color: #ffffff; color: #000;
            font-weight: bold; text-align: center; vertical-align: middle;
            line-height: 1.25; height: 68px; z-index: 10; border-bottom: 2px solid #475569;
        }}
        .red-txt {{ color: #e11d48; font-weight: bold; font-size: 10.5px; }}
        .sticky-col-1 {{ position: sticky; left: 0px; z-index: 5; background: #fff; }}
        .sticky-col-2 {{ position: sticky; left: 45px; z-index: 5; background: #fff; }}
        .sticky-col-3 {{ position: sticky; left: 110px; z-index: 5; background: #fff; }}
        .sticky-col-4 {{ position: sticky; left: 245px; z-index: 5; background: #fff; border-right: 2px solid #64748b; }}
        thead th.sticky-col-1, thead th.sticky-col-2, thead th.sticky-col-3, thead th.sticky-col-4 {{ z-index: 15; }}

        .w-no {{ width: 45px; min-width: 45px; text-align: center; }}
        .w-center {{ width: 65px; min-width: 65px; text-align: center; }}
        .w-code {{ width: 135px; min-width: 135px; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .w-name {{ width: 200px; min-width: 200px; max-width: 200px; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .w-num {{ width: 72px; min-width: 72px; text-align: right; font-variant-numeric: tabular-nums; }}
        .w-date {{ width: 95px; min-width: 95px; text-align: center; }}
        .w-prev-dec {{ width: 160px; min-width: 160px; max-width: 160px; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .w-dec {{ width: 220px; min-width: 220px; }}

        .bg-cyan {{ background-color: #67e8f9 !important; }}
        .bg-flag-neg {{ background-color: #fecdd3 !important; color: #be123c; font-weight: bold; }}
        .bg-green {{ background-color: #a3e635 !important; }}
        .bg-pink-rate {{ background-color: #fbcfe8 !important; color: #be123c; font-weight: bold; }}
        .bg-yellow-ro {{ background-color: #fef08a !important; color: #1e293b; }}
        .bg-yellow-edit {{ background-color: #ffff00 !important; }}
        .bg-sales-hit {{ background-color: #fed7aa !important; font-weight: bold; color: #c2410c; }}
        .bg-hist-ki {{ background-color: #f1f5f9; color: #475569; }}
        .bg-zure-diff {{ background-color: #ffe4e6 !important; color: #be123c; font-weight: bold; }}

        tbody tr:hover td {{ filter: brightness(0.95); }}
        .click-link {{ color: #0284c7; font-weight: bold; text-decoration: underline; cursor: pointer; display: block; width: 100%; }}
        .cell-input {{ width: 100%; height: 23px; border: 1px solid #94a3b8; border-radius: 2px; padding: 1px 4px; font-size: 12px; background: rgba(255,255,255,0.9); }}
        .cell-input-point {{ width: 100%; height: 23px; border: 1px solid #3b82f6; border-radius: 2px; padding: 1px 4px; font-size: 12px; font-weight: bold; text-align: right; background: #eff6ff; }}

        .modal-overlay {{
            display: none; position: fixed; top: 0; left: 0; width: 100%; height: 100%;
            background: rgba(15, 23, 42, 0.7); z-index: 9999; justify-content: center; align-items: center;
        }}
        .modal-window {{
            background: #fff; width: 99%; max-width: 1750px; height: 96vh; border-radius: 10px;
            box-shadow: 0 25px 40px -5px rgba(0,0,0,0.45); display: flex; flex-direction: column; overflow: hidden; border: 3px solid #1e40af;
        }}
        .modal-header {{ background: #1e3a8a; color: #fff; padding: 16px 24px; display: flex; justify-content: space-between; align-items: center; }}
        .modal-title {{ font-size: 20px; font-weight: bold; }}
        .modal-close {{ background: #ef4444; color: #fff; border: none; border-radius: 6px; padding: 8px 22px; font-size: 15px; font-weight: bold; cursor: pointer; }}
        .modal-body {{ padding: 22px; overflow-y: auto; display: flex; flex-direction: column; gap: 18px; flex: 1; }}
        .summary-bar {{ display: flex; gap: 24px; flex-wrap: wrap; background: #f1f5f9; padding: 14px 20px; border-radius: 8px; border-left: 6px solid #2563eb; font-size: 15.5px; }}
        .summary-item {{ font-weight: bold; color: #1e293b; }}
        .summary-item span {{ color: #0284c7; font-size: 18px; margin-left: 5px; }}
        .quick-edit-box {{ background: #fffbeb; border: 2px solid #fde68a; padding: 14px 18px; border-radius: 8px; display: flex; flex-direction: column; gap: 12px; }}
        .stamp-btn {{ background: #e2e8f0; border: 1px solid #94a3b8; border-radius: 5px; padding: 5px 12px; font-size: 14px; cursor: pointer; font-weight: bold; }}
        .modal-grid {{ display: grid; grid-template-columns: 3.5fr 1.5fr; gap: 20px; }}
        .sub-table {{ width: 100%; border-collapse: collapse; }}
        .sub-table th {{ background: #e2e8f0; color: #1e293b; padding: 10px 12px; font-size: 14.5px; height: auto; position: static; border: 1px solid #cbd5e1; }}
        .sub-table td {{ padding: 10px 12px; font-size: 14.5px; height: auto; border: 1px solid #cbd5e1; }}
        .sec-title {{ font-size: 16.5px; font-weight: bold; color: #1e3a8a; margin-bottom: 8px; border-bottom: 2px solid #93c5fd; padding-bottom: 4px; }}
    </style>
    </head>
    <body>
    <div class="toolbar">
        <div class="toolbar-info">表示件数: <span id="row-count">0</span> 件 ｜ 💡 発注点入力後はEnterキーで即座に下の行へ移動します</div>
        <button class="btn-download" onclick="exportEditedCSV()">📥 編集内容を反映したCSV出力</button>
    </div>
    <div class="table-wrap">
        <table>
            <thead>
                <tr>
                    <th class="w-no sticky-col-1">No</th>
                    <th class="w-center sticky-col-2">商品<br>センター</th>
                    <th class="w-code sticky-col-3">品番</th>
                    <th class="w-name sticky-col-4">商品区分</th>
                    <th class="w-num bg-hist-ki">73期<br>実績</th>
                    <th class="w-num bg-hist-ki">74期<br>実績</th>
                    <th class="w-num bg-hist-ki">75期<br>実績</th>
                    <th class="w-num bg-hist-ki">76期<br>実績</th>
                    <th class="w-num" style="background:#e0f2fe;">1課発注点<br><span style="color:#0369a1; font-size:10px;">✏️手動修正可</span></th>
                    <th class="w-num" style="background:#e0f2fe;">2課発注点<br><span style="color:#0369a1; font-size:10px;">✏️手動修正可</span></th>
                    <th class="w-num bg-cyan">在庫数<br><span class="red-txt">ロジザード</span></th>
                    <th class="w-num bg-cyan">入荷予定<br>数<br><span class="red-txt">ロジザード</span></th>
                    <th class="w-date bg-cyan">入荷予定<br>日<br><span class="red-txt">ロジザード</span></th>
                    <th class="w-num bg-cyan">確保数<br><span class="red-txt">ロジザード<br>(在庫確保)</span></th>
                    <th class="w-num bg-cyan">在庫確保<br><span class="red-txt">(弥生)</span></th>
                    <th class="w-num bg-cyan">在庫確保<br><span style="color:#b91c1c; font-weight:bold;">丸藤</span></th>
                    <th class="w-num bg-cyan">確保数の<br>ズレ</th>
                    <th class="w-num bg-cyan">入荷後の<br>フリー在庫</th>
                    <th class="w-num">発注検討<br>フラグ<br><span class="red-txt">マイナスの<br>場合発注検討</span></th>
                    <th class="w-num">今月売れ数<br><span style="color:#c2410c; font-weight:bold;">🔥実データ</span></th>
                    <th class="w-num">在庫回転<br>率</th>
                    <th class="w-num">リンベル<br>確保数</th>
                    <th class="w-prev-dec bg-yellow-ro">{prev_title_html}</th>
                    <th class="w-dec">決定事項</th>
                    <th class="w-date">発注日</th>
                </tr>
            </thead>
            <tbody id="table-body"></tbody>
        </table>
    </div>

    <div class="modal-overlay" id="detail-modal" onclick="closeModalOnBg(event)">
        <div class="modal-window">
            <div class="modal-header">
                <div class="modal-title" id="modal-title-text">📦 品番詳細・確保データ</div>
                <button class="modal-close" onclick="closeModal()">✕ 閉じる</button>
            </div>
            <div class="modal-body">
                <div class="summary-bar" id="modal-summary"></div>
                <div class="quick-edit-box">
                    <div style="font-size:14.5px; font-weight:bold; color:#92400e; display:flex; justify-content:space-between;">
                        <span>✏️ この場で決定事項・入荷予定日を入力（表にもリアルタイム連動）</span>
                        <span>
                            スタンプ:
                            <button class="stamp-btn" onclick="addStamp('様子見')">＋様子見</button>
                            <button class="stamp-btn" onclick="addStamp('発注済')">＋発注済</button>
                            <button class="stamp-btn" onclick="addStamp('確認中')">＋確認中</button>
                            <button class="stamp-btn" onclick="addStamp('廃盤予定')">＋廃盤予定</button>
                        </span>
                    </div>
                    <div style="display:flex; gap:12px; align-items:center;">
                        <label style="font-size:14px; font-weight:bold;">前回:</label>
                        <input type="text" id="modal-view-prev" class="cell-input" style="width:200px; height:32px; background:#fef08a; font-size:14px;" readonly>
                        <label style="font-size:14px; font-weight:bold;">入荷予定日:</label>
                        <input type="text" id="modal-input-ndate" class="cell-input" style="width:140px; height:32px; font-size:14px;" oninput="syncFromModal()">
                        <label style="font-size:14px; font-weight:bold;">決定事項:</label>
                        <input type="text" id="modal-input-dec" class="cell-input" style="flex:1; height:32px; font-size:14px;" oninput="syncFromModal()">
                        <label style="font-size:14px; font-weight:bold;">発注日:</label>
                        <input type="text" id="modal-input-odate" class="cell-input" style="width:130px; height:32px; font-size:14px;" oninput="syncFromModal()">
                    </div>
                </div>
                <div class="modal-grid">
                    <div>
                        <div class="sec-title" id="modal-sub-title">🏢 得意先別・弥生確保リスト（受注日の最新順）</div>
                        <div id="modal-yayoi-area"></div>
                    </div>
                    <div>
                        <div class="sec-title">📜 過去の会議決定履歴（最新順）</div>
                        <div id="modal-hist-area"></div>
                    </div>
                </div>
            </div>
        </div>
    </div>

    <script>
        const tableData = {records_json};
        const yayoiMap = {yayoi_json};
        const marufujiMap = {marufuji_json};
        const histMap = {hist_json};
        let currentModalIdx = null;

        const savedEdits = JSON.parse(localStorage.getItem("hattyu_edits_v6") || "{{}}");
        tableData.forEach(r => {{
            if (savedEdits[r["品番"]]) {{
                r["1課発注点"] = savedEdits[r["品番"]].p1 ?? r["1課発注点"];
                r["2課発注点"] = savedEdits[r["品番"]].p2 ?? r["2課発注点"];
                r["入荷予定日_ロジザード"] = savedEdits[r["品番"]].ndate ?? r["入荷予定日_ロジザード"];
                r["決定事項"] = savedEdits[r["品番"]].dec ?? r["決定事項"];
                r["発注日"] = savedEdits[r["品番"]].odate ?? r["発注日"];
                r["発注検討フラグ"] = Number(r["入荷後のフリー在庫"] || 0) - (Number(r["1課発注点"] || 0) + Number(r["2課発注点"] || 0));
            }}
            r["確保数のズレ"] = Number(r["確保数_ロジザード"] || 0) - Number(r["在庫確保_弥生"] || 0) - Number(r["在庫確保_丸藤"] || 0);
        }});

        function fmtNum(val) {{
            if (val === "" || val === null || val === undefined) return "";
            const n = Number(val);
            return isNaN(n) ? val : n.toLocaleString("ja-JP");
        }}

        function handleKeyDown(event, colName, currentRowIdx) {{
            if (event.key === "Enter") {{
                event.preventDefault();
                const nextRowIdx = currentRowIdx + 1;
                if (nextRowIdx < tableData.length) {{
                    const nextInput = document.getElementById(`${{colName}}-${{nextRowIdx}}`);
                    if (nextInput) {{
                        nextInput.focus();
                        if (nextInput.select) nextInput.select();
                    }}
                }}
            }}
        }}

        function recalculateRow(idx) {{
            const r = tableData[idx];
            const p1 = Number(r["1課発注点"] || 0);
            const p2 = Number(r["2課発注点"] || 0);
            const free = Number(r["入荷後のフリー在庫"] || 0);
            r["発注検討フラグ"] = free - (p1 + p2);

            const flagCell = document.getElementById(`flag-${{idx}}`);
            if (flagCell) {{
                flagCell.textContent = fmtNum(r["発注検討フラグ"]);
                if (r["発注検討フラグ"] < 0) {{
                    flagCell.className = "w-num bg-flag-neg";
                }} else {{
                    flagCell.className = "w-num bg-cyan";
                }}
            }}
            saveToStorage(r["品番"], r);
        }}

        function renderTable() {{
            const tbody = document.getElementById("table-body");
            document.getElementById("row-count").textContent = tableData.length.toLocaleString();
            let html = "";
            tableData.forEach((r, idx) => {{
                const flagVal = Number(r["発注検討フラグ"]);
                const flagCls = (flagVal < 0) ? "bg-flag-neg" : "bg-cyan";
                const rateVal = Number(r["在庫回転率"]);
                const rateCls = (rateVal <= 5) ? "bg-pink-rate" : "bg-green";
                const salesVal = Number(r["今月売れ数"] || 0);
                const salesCls = (salesVal >= 1) ? "bg-sales-hit" : "";
                const zureVal = Number(r["確保数のズレ"] || 0);
                const zureCls = (zureVal !== 0) ? "bg-zure-diff" : "bg-cyan";

                html += `<tr>
                    <td class="w-no sticky-col-1">${{r["No"]}}</td>
                    <td class="w-center sticky-col-2">${{r["商品センター"]}}</td>
                    <td class="w-code sticky-col-3" title="${{r["品番"]}}"><span class="click-link" onclick="openModal(${{idx}}, 'all')">${{r["品番"]}}</span></td>
                    <td class="w-name sticky-col-4" title="${{r["商品区分"]}}">${{r["商品区分"]}}</td>
                    <td class="w-num bg-hist-ki">${{fmtNum(r["73期"])}}</td>
                    <td class="w-num bg-hist-ki">${{fmtNum(r["74期"])}}</td>
                    <td class="w-num bg-hist-ki">${{fmtNum(r["75期"])}}</td>
                    <td class="w-num bg-hist-ki">${{fmtNum(r["76期"])}}</td>
                    <td class="w-num"><input type="number" class="cell-input-point" id="point1-${{idx}}" value="${{r["1課発注点"]}}" oninput="tableData[${{idx}}]['1課発注点']=Number(this.value); recalculateRow(${{idx}})" onkeydown="handleKeyDown(event, 'point1', ${{idx}})"></td>
                    <td class="w-num"><input type="number" class="cell-input-point" id="point2-${{idx}}" value="${{r["2課発注点"]}}" oninput="tableData[${{idx}}]['2課発注点']=Number(this.value); recalculateRow(${{idx}})" onkeydown="handleKeyDown(event, 'point2', ${{idx}})"></td>
                    <td class="w-num bg-cyan">${{fmtNum(r["在庫数_ロジザード"])}}</td>
                    <td class="w-num bg-cyan">${{fmtNum(r["入荷予定数_ロジザード"])}}</td>
                    <td class="w-date bg-cyan"><input type="text" class="cell-input" id="ndate-${{idx}}" value="${{r["入荷予定日_ロジザード"] || ""}}" oninput="updateCell(${{idx}}, '入荷予定日_ロジザード", this.value)" onkeydown="handleKeyDown(event, 'ndate', ${{idx}})"></td>
                    <td class="w-num bg-cyan">${{fmtNum(r["確保数_ロジザード"])}}</td>
                    <td class="w-num bg-cyan"><span class="click-link" onclick="openModal(${{idx}}, 'yayoi')">${{fmtNum(r["在庫確保_弥生"])}}</span></td>
                    <td class="w-num bg-cyan"><span class="click-link" style="color:#b91c1c;" onclick="openModal(${{idx}}, 'marufuji')">${{fmtNum(r["在庫確保_丸藤"])}}</span></td>
                    <td class="w-num ${{zureCls}}" id="zure-${{idx}}">${{fmtNum(r["確保数のズレ"])}}</td>
                    <td class="w-num bg-cyan">${{fmtNum(r["入荷後のフリー在庫"])}}</td>
                    <td class="w-num ${{flagCls}}" id="flag-${{idx}}">${{fmtNum(r["発注検討フラグ"])}}</td>
                    <td class="w-num ${{salesCls}}">${{salesVal > 0 ? "🔥 " + fmtNum(salesVal) : "0"}}</td>
                    <td class="w-num ${{rateCls}}">${{fmtNum(r["在庫回転率"]) + "%"}}</td>
                    <td class="w-num">${{fmtNum(r["リンベル確保数"])}}</td>
                    <td class="w-prev-dec bg-yellow-ro" title="${{r["前回決定事項"] || ""}}">${{r["前回決定事項"] || ""}}</td>
                    <td class="w-dec bg-yellow-edit"><input type="text" class="cell-input" id="dec-${{idx}}" value="${{r["決定事項"] || ""}}" oninput="updateCell(${{idx}}, '決定事項", this.value)" onkeydown="handleKeyDown(event, 'dec', ${{idx}})"></td>
                    <td class="w-date bg-yellow-edit"><input type="text" class="cell-input" id="odate-${{idx}}" value="${{r["発注日"] || ""}}" oninput="updateCell(${{idx}}, '発注日', this.value)" onkeydown="handleKeyDown(event, 'odate', ${{idx}})"></td>
                </tr>`;
            }});
            tbody.innerHTML = html;
        }}

        function saveToStorage(hinban, rowObj) {{
            const store = JSON.parse(localStorage.getItem("hattyu_edits_v6") || "{{}}");
            store[hinban] = {{
                p1: rowObj["1課発注点"],
                p2: rowObj["2課発注点"],
                ndate: rowObj["入荷予定日_ロジザード"],
                dec: rowObj["決定事項"],
                odate: rowObj["発注日"]
            }};
            localStorage.setItem("hattyu_edits_v6", JSON.stringify(store));
        }}

        function updateCell(idx, field, val) {{
            tableData[idx][field] = val;
            saveToStorage(tableData[idx]["品番"], tableData[idx]);
        }}

        function openModal(idx, mode) {{
            currentModalIdx = idx;
            const r = tableData[idx];
            const code = r["品番"];
            const codeNorm = r["品番_正規化"] || code;

            document.getElementById("modal-title-text").innerHTML = `📦 品番: 【${{code}}】 ${{r["商品区分"]}} の詳細`;
            document.getElementById("modal-summary").innerHTML = `
                <div class="summary-item">センター: <span>${{r["商品センター"]}}</span></div>
                <div class="summary-item">1課発注点: <span>${{fmtNum(r["1課発注点"])}}</span></div>
                <div class="summary-item">2課発注点: <span>${{fmtNum(r["2課発注点"])}}</span></div>
                <div class="summary-item">ロジ在庫: <span>${{fmtNum(r["在庫数_ロジザード"])}}</span></div>
                <div class="summary-item">弥生確保: <span>${{fmtNum(r["在庫確保_弥生"])}}</span></div>
                <div class="summary-item">丸藤確保: <span style="color:#b91c1c;">${{fmtNum(r["在庫確保_丸藤"])}}</span></div>
                <div class="summary-item">確保数のズレ: <span style="color:#e11d48;">${{fmtNum(r["確保数のズレ"])}}</span></div>
                <div class="summary-item">フリー在庫: <span>${{fmtNum(r["入荷後のフリー在庫"])}}</span></div>
                <div class="summary-item">検討フラグ: <span style="color:#e11d48;">${{fmtNum(r["発注検討フラグ"])}}</span></div>
            `;
            document.getElementById("modal-view-prev").value = r["前回決定事項"] || "";
            document.getElementById("modal-input-ndate").value = r["入荷予定日_ロジザード"] || "";
            document.getElementById("modal-input-dec").value = r["決定事項"] || "";
            document.getElementById("modal-input-odate").value = r["発注日"] || "";

            let targetList = [];
            if (mode === 'marufuji') {{
                document.getElementById("modal-sub-title").innerHTML = "🏢 【丸藤専用】確保リスト（丸藤の受注のみ表示）";
                targetList = marufujiMap[codeNorm] || [];
            }} else {{
                document.getElementById("modal-sub-title").innerHTML = "🏢 得意先別・弥生確保リスト（受注日の最新順）";
                targetList = yayoiMap[codeNorm] || [];
            }}

            if (targetList.length === 0) {{
                document.getElementById("modal-yayoi-area").innerHTML = `<div style="padding:16px; font-size:14.5px; color:#64748b; background:#f8fafc;">該当する受注データはありません。</div>`;
            }} else {{
                let yHtml = `<table class="sub-table"><thead><tr><th>受注番号</th><th>得意先名</th><th>数量</th><th>受注日</th><th>納入期日</th><th>倉庫名</th><th>担当者</th><th>備考</th></tr></thead><tbody>`;
                targetList.forEach(y => {{
                    yHtml += `<tr>
                        <td style="text-align:center; font-weight:bold;">${{y["受注番号"] || ""}}</td>
                        <td style="font-weight:bold;">${{y["得意先名"]}}</td>
                        <td style="text-align:right; font-weight:bold; color:#0284c7; font-size:15.5px;">${{fmtNum(y["数量"])}}</td>
                        <td style="text-align:center;">${{y["受注日"]}}</td>
                        <td style="text-align:center;">${{y["納入期日"]}}</td>
                        <td style="text-align:center; font-weight:bold;">${{y["倉庫名"] || ""}}</td>
                        <td style="text-align:center;">${{y["担当者名"] || ""}}</td>
                        <td>${{y["備考"]}}</td>
                    </tr>`;
                }});
                yHtml += `</tbody></table>`;
                document.getElementById("modal-yayoi-area").innerHTML = yHtml;
            }}

            const hList = histMap[code] || [];
            if (hList.length === 0) {{
                document.getElementById("modal-hist-area").innerHTML = `<div style="padding:16px; font-size:14.5px; color:#64748b; background:#f8fafc;">過去の会議決定履歴はありません。</div>`;
            }} else {{
                let hHtml = `<table class="sub-table"><thead><tr><th style="width:110px;">会議日</th><th>決定事項</th></tr></thead><tbody>`;
                hList.forEach(h => {{
                    hHtml += `<tr><td style="text-align:center; font-weight:bold;">${{h["会議日"]}}</td><td style="font-weight:bold;">${{h["決定事項"]}}</td></tr>`;
                }});
                hHtml += `</tbody></table>`;
                document.getElementById("modal-hist-area").innerHTML = hHtml;
            }}
            document.getElementById("detail-modal").style.display = "flex";
        }}

        function syncFromModal() {{
            if (currentModalIdx === null) return;
            const ndate = document.getElementById("modal-input-ndate").value;
            const dec = document.getElementById("modal-input-dec").value;
            const odate = document.getElementById("modal-input-odate").value;
            tableData[currentModalIdx]["入荷予定日_ロジザード"] = ndate;
            tableData[currentModalIdx]["決定事項"] = dec;
            tableData[currentModalIdx]["発注日"] = odate;
            document.getElementById(`ndate-${{currentModalIdx}}`).value = ndate;
            document.getElementById(`dec-${{currentModalIdx}}`).value = dec;
            document.getElementById(`odate-${{currentModalIdx}}`).value = odate;
            saveToStorage(tableData[currentModalIdx]["品番"], tableData[currentModalIdx]);
        }}

        function addStamp(text) {{
            const decInput = document.getElementById("modal-input-dec");
            const curr = decInput.value.trim();
            decInput.value = curr ? `${{curr}} / ${{text}}` : text;
            syncFromModal();
        }}

        function closeModal() {{ document.getElementById("detail-modal").style.display = "none"; currentModalIdx = null; }}
        function closeModalOnBg(e) {{ if (e.target.id === "detail-modal") closeModal(); }}

        function exportEditedCSV() {{
            const cols = ["No", "商品センター", "品番", "商品区分", "73期", "74期", "75期", "76期", "1課発注点", "2課発注点", "在庫数_ロジザード", "入荷予定数_ロジザード", "入荷予定日_ロジザード", "確保数_ロジザード", "在庫確保_弥生", "在庫確保_丸藤", "確保数のズレ", "入荷後のフリー在庫", "発注検討フラグ", "今月売れ数", "在庫回転率", "リンベル確保数", "前回決定事項", "決定事項", "発注日"];
            let csv = "\\uFEFF" + cols.join(",") + "\\n";
            tableData.forEach(r => {{
                csv += cols.map(c => `"${{String(r[c] ?? "").replace(/"/g, '""')}}"`).join(",") + "\\n";
            }});
            const blob = new Blob([csv], {{ type: "text/csv;charset=utf-8;" }});
            const url = URL.createObjectURL(blob);
            const a = document.createElement("a");
            a.href = url;
            a.download = `発注検討表_出力_${{new Date().toISOString().slice(0,10)}}.csv`;
            a.click();
        }}
        renderTable();
    </script>
    </body>
    </html>
    """
    components.html(html_tab1, height=680, scrolling=False)

with tab2:
    if df_yayoi.empty:
        st.warning("確保リストのデータがありません。")
    else:
        comp_totals = df_yayoi.groupby("得意先名")["数量"].sum().sort_values(ascending=False)
        comp_opts = ["すべての会社を表示"] + [f"{c} （確保合計: {v:,} 個）" for c, v in comp_totals.items()]
        c1, c2 = st.columns(2)
        sel_comp = c1.selectbox("🏢 会社名（得意先）をプルダウンで選択", comp_opts)
        kw_comp = c2.text_input("🔍 受注番号・得意先名・品番・備考で検索", placeholder="例: 00000001 / オバタ / 丸藤 / AHL-2203")

        yv = df_yayoi.copy()
        if sel_comp != "すべての会社を表示":
            cname = sel_comp.split(" （確保合計:")[0]
            yv = yv[yv["得意先名"] == cname]
        if kw_comp:
            yv = yv[
                yv["受注番号"].astype(str).str.contains(kw_comp, case=False, na=False) |
                yv["得意先名"].str.contains(kw_comp, case=False, na=False) |
                yv["商品コード"].str.contains(kw_comp, case=False, na=False) |
                yv["備考"].str.contains(kw_comp, case=False, na=False)
            ]

        yv_full = pd.merge(
            yv,
            m_df[["品番_正規化", "商品区分", "在庫数_ロジザード", "入荷予定日_ロジザード", "入荷後のフリー在庫", "発注検討フラグ", "決定事項"]],
            on="品番_正規化", how="left"
        ).fillna("")

        slips_summary_list = []
        slip_details_map = {}

        for ono, grp in yv_full.groupby("受注番号", sort=False):
            ono_str = str(ono)
            first_r = grp.iloc[0]
            codes_preview = " / ".join(list(grp["商品コード"].unique())[:4])
            if len(grp["商品コード"].unique()) > 4:
                codes_preview += " 他"

            slips_summary_list.append({
                "受注番号": ono_str,
                "受注日": clean_str(first_r["受注日"]),
                "納入期日": clean_str(first_r["納入期日"]),
                "得意先名": clean_str(first_r["得意先名"]),
                "担当者名": clean_str(first_r.get("担当者名", "")),
                "登録商品数": len(grp),
                "受注残数合計": int(grp["数量"].sum()),
                "商品内訳プレビュー": codes_preview,
                "備考": clean_str(first_r.get("備考", ""))
            })

            d_list = []
            for _, dr in grp.iterrows():
                p_name = clean_str(dr.get("商品区分", "")) or clean_str(dr.get("商品名/摘要", "")) or item_name_map.get(clean_str(dr["商品コード"]), "")
                d_list.append({
                    "商品コード": clean_str(dr["商品コード"]),
                    "商品名": p_name,
                    "数量": int(dr["数量"]),
                    "倉庫名": clean_str(dr.get("倉庫名", "")),
                    "納入期日": clean_str(dr.get("納入期日", "")),
                    "備考": clean_str(dr.get("備考", "")),
                    "ロジ在庫": int(to_num(dr.get("在庫数_ロジザード", 0))),
                    "フリー在庫": int(to_num(dr.get("入荷後のフリー在庫", 0))),
                    "入荷予定日": clean_str(dr.get("入荷予定日_ロジザード", "")),
                    "決定事項": clean_str(dr.get("決定事項", ""))
                })
            slip_details_map[ono_str] = d_list

        slips_summary_list.sort(key=lambda x: (x["受注日"], x["受注番号"]), reverse=True)
        slips_json = json.dumps(slips_summary_list, ensure_ascii=False)
        slip_details_json = json.dumps(slip_details_map, ensure_ascii=False)

        html_tab2 = f"""
        <!DOCTYPE html>
        <html>
        <head>
        <meta charset="utf-8">
        <style>
            * {{ box-sizing: border-box; font-family: "Meiryo", "Hiragino Kaku Gothic ProN", sans-serif; }}
            body {{ margin: 0; padding: 0; background: #fff; }}
            .toolbar {{
                display: flex; align-items: center; justify-content: space-between;
                background: #f8fafc; padding: 8px 14px; border: 1px solid #cbd5e1;
                border-bottom: none; border-radius: 6px 6px 0 0;
            }}
            .toolbar-info {{ font-size: 13.5px; color: #1e293b; font-weight: bold; }}
            .toolbar-hint {{ font-size: 12.5px; color: #0369a1; background: #e0f2fe; padding: 4px 12px; border-radius: 4px; font-weight: bold; }}
            .table-wrap {{ width: 100%; height: 580px; overflow: auto; border: 1px solid #94a3b8; }}
            table {{ border-collapse: separate; border-spacing: 0; width: 100%; }}
            th, td {{ border-right: 1px solid #cbd5e1; border-bottom: 1px solid #cbd5e1; padding: 7px 10px; font-size: 13.5px; }}
            thead th {{
                position: sticky; top: 0; background: #1e3a8a; color: #fff;
                font-weight: bold; text-align: center; height: 40px; z-index: 10;
            }}
            tbody tr:nth-child(even) {{ background: #f8fafc; }}
            tbody tr:hover {{ background: #e0f2fe; cursor: pointer; }}
            .click-link {{ color: #0284c7; font-weight: bold; text-decoration: underline; cursor: pointer; }}
            .click-link:hover {{ color: #1d4ed8; }}

            .modal-overlay {{
                display: none; position: fixed; top: 0; left: 0; width: 100%; height: 100%;
                background: rgba(15, 23, 42, 0.7); z-index: 9999; justify-content: center; align-items: center;
            }}
            .modal-window {{
                background: #fff; width: 99%; max-width: 1750px; height: 96vh; border-radius: 10px;
                box-shadow: 0 25px 40px -5px rgba(0,0,0,0.45); display: flex; flex-direction: column; overflow: hidden; border: 3px solid #1e40af;
            }}
            .modal-header {{ background: #1e3a8a; color: #fff; padding: 16px 24px; display: flex; justify-content: space-between; align-items: center; }}
            .modal-title {{ font-size: 20px; font-weight: bold; }}
            .modal-close {{ background: #ef4444; color: #fff; border: none; border-radius: 6px; padding: 8px 22px; font-size: 15px; font-weight: bold; cursor: pointer; }}
            .modal-body {{ padding: 22px; overflow-y: auto; display: flex; flex-direction: column; gap: 18px; flex: 1; }}
            .summary-bar {{ display: flex; gap: 24px; flex-wrap: wrap; background: #f1f5f9; padding: 14px 20px; border-radius: 8px; border-left: 6px solid #2563eb; font-size: 15.5px; }}
            .summary-item {{ font-weight: bold; color: #1e293b; }}
            .summary-item span {{ color: #0284c7; font-size: 18px; margin-left: 5px; }}
            .sub-table {{ width: 100%; border-collapse: collapse; }}
            .sub-table th {{ background: #e2e8f0; color: #1e293b; padding: 10px 12px; font-size: 14.5px; border: 1px solid #cbd5e1; position: static; height: auto; }}
            .sub-table td {{ padding: 10px 12px; font-size: 14.5px; border: 1px solid #cbd5e1; }}
            .sec-title {{ font-size: 16.5px; font-weight: bold; color: #1e3a8a; border-bottom: 2px solid #93c5fd; padding-bottom: 4px; }}
        </style>
        </head>
        <body>
        <div class="toolbar">
            <div class="toolbar-info">受注伝票件数: <span id="slip-count">0</span> 件（受注番号ごとに集約）</div>
            <div class="toolbar-hint">💡 行をクリックすると、特大ポップアップで登録内容が開きます！</div>
        </div>
        <div class="table-wrap">
            <table>
                <thead>
                    <tr>
                        <th style="width:130px;">受注番号</th>
                        <th style="width:105px;">受注日</th>
                        <th style="width:105px;">納入期日</th>
                        <th style="width:260px;">得意先名（会社名）</th>
                        <th style="width:85px;">担当者</th>
                        <th style="width:85px;">商品行数</th>
                        <th style="width:110px;">受注残数 合計</th>
                        <th>登録商品コード内訳</th>
                        <th style="width:180px;">備考 / 摘要</th>
                    </tr>
                </thead>
                <tbody id="slip-tbody"></tbody>
            </table>
        </div>

        <div class="modal-overlay" id="slip-modal" onclick="closeSlipModalOnBg(event)">
            <div class="modal-window">
                <div class="modal-header">
                    <div class="modal-title" id="slip-modal-title">📦 受注伝票 登録内容詳細</div>
                    <button class="modal-close" onclick="closeSlipModal()">✕ 閉じる</button>
                </div>
                <div class="modal-body">
                    <div class="summary-bar" id="slip-modal-summary"></div>
                    <div class="sec-title">📋 登録されている商品明細一覧</div>
                    <div id="slip-modal-table-area"></div>
                </div>
            </div>
        </div>

        <script>
            const slipsData = {slips_json};
            const slipDetailsMap = {slip_details_json};

            function fmtNum(val) {{
                const n = Number(val);
                return isNaN(n) ? val : n.toLocaleString("ja-JP");
            }}

            function renderSlips() {{
                document.getElementById("slip-count").textContent = slipsData.length.toLocaleString();
                const tbody = document.getElementById("slip-tbody");
                let html = "";
                slipsData.forEach((s, idx) => {{
                    html += `<tr onclick="openSlipModal(${{idx}})">
                        <td style="text-align:center;"><span class="click-link">${{s["受注番号"]}}</span></td>
                        <td style="text-align:center;">${{s["受注日"]}}</td>
                        <td style="text-align:center;">${{s["納入期日"]}}</td>
                        <td><span class="click-link">${{s["得意先名"]}}</span></td>
                        <td style="text-align:center;">${{s["担当者名"]}}</td>
                        <td style="text-align:right; font-weight:bold;">${{fmtNum(s["登録商品数"])}} 点</td>
                        <td style="text-align:right; font-weight:bold; color:#0284c7;">${{fmtNum(s["受注残数合計"])}} 個</td>
                        <td style="color:#334155;">${{s["商品内訳プレビュー"]}}</td>
                        <td>${{s["備考"]}}</td>
                    </tr>`;
                }});
                tbody.innerHTML = html;
            }}

            function openSlipModal(idx) {{
                const s = slipsData[idx];
                const ono = s["受注番号"];
                const details = slipDetailsMap[ono] || [];

                document.getElementById("slip-modal-title").innerHTML = `🏢 受注番号: 【${{ono}}】 ${{s["得意先名"]}} の登録内容詳細`;
                document.getElementById("slip-modal-summary").innerHTML = `
                    <div class="summary-item">得意先名: <span>${{s["得意先名"]}}</span></div>
                    <div class="summary-item">受注番号: <span>${{ono}}</span></div>
                    <div class="summary-item">受注日: <span>${{s["受注日"]}}</span></div>
                    <div class="summary-item">納入期日: <span>${{s["納入期日"]}}</span></div>
                    <div class="summary-item">担当者: <span>${{s["担当者名"] || "—"}}</span></div>
                    <div class="summary-item">明細行数: <span>${{fmtNum(s["登録商品数"])}} 行</span></div>
                    <div class="summary-item">受注残数 合計: <span style="color:#e11d48;">${{fmtNum(s["受注残数合計"])}} 個</span></div>
                `;

                let tHtml = `<table class="sub-table"><thead><tr>
                    <th style="width:55px;">No</th>
                    <th style="width:150px;">商品コード</th>
                    <th>商品名</th>
                    <th style="width:110px;">受注残数</th>
                    <th style="width:100px;">倉庫名</th>
                    <th style="width:115px;">納入期日</th>
                    <th style="width:105px;">ロジ在庫</th>
                    <th style="width:105px;">フリー在庫</th>
                    <th style="width:115px;">入荷予定日</th>
                    <th>備考 / 決定事項</th>
                </tr></thead><tbody>`;

                details.forEach((d, i) => {{
                    tHtml += `<tr>
                        <td style="text-align:center;">${{i + 1}}</td>
                        <td style="font-weight:bold; color:#0284c7;">${{d["商品コード"]}}</td>
                        <td style="font-weight:bold;">${{d["商品名"]}}</td>
                        <td style="text-align:right; font-weight:bold; color:#be123c; background:#fff1f2; font-size:15.5px;">${{fmtNum(d["数量"])}}</td>
                        <td style="text-align:center; font-weight:bold;">${{d["倉庫名"]}}</td>
                        <td style="text-align:center;">${{d["納入期日"]}}</td>
                        <td style="text-align:right; background:#ecfeff;">${{fmtNum(d["ロジ在庫"])}}</td>
                        <td style="text-align:right; background:#ecfeff;">${{fmtNum(d["フリー在庫"])}}</td>
                        <td style="text-align:center;">${{d["入荷予定日"]}}</td>
                        <td>${{d["備考"]}} ${{d["決定事項"] ? "（決定: " + d["決定事項"] + "）" : ""}}</td>
                    </tr>`;
                }});
                tHtml += `</tbody></table>`;
                document.getElementById("slip-modal-table-area").innerHTML = tHtml;
                document.getElementById("slip-modal").style.display = "flex";
            }}

            function closeSlipModal() {{ document.getElementById("slip-modal").style.display = "none"; }}
            function closeSlipModalOnBg(e) {{ if (e.target.id === "slip-modal") closeSlipModal(); }}

            renderSlips();
        </script>
        </body>
        </html>
        """
        components.html(html_tab2, height=650, scrolling=False)

with tab3:
    max_seq = 0
    if os.path.exists(MANUAL_YAYOI_CSV):
        try:
            df_m_check = pd.read_csv(MANUAL_YAYOI_CSV)
            for ono in df_m_check.get("受注番号", []):
                s_ono = clean_str(ono).replace("WEB-", "")
                if s_ono.isdigit():
                    max_seq = max(max_seq, int(s_ono))
        except Exception:
            pass
    next_order_no_str = f"{max_seq + 1:08d}"

    comp_select_list = [""] + [f"{c['name']}" + (f" （{c['code']}）" if c["code"] else "") for c in master_comps]
    item_code_list = [""] + [it["code"] for it in master_items]

    def create_blank_slip(order_no_val):
        return {
            "order_no": order_no_val,
            "order_date": datetime.date.today(),
            "delivery_date": datetime.date.today() + datetime.timedelta(days=5),
            "comp_label": "",
            "staff": "",
            "summary": "",
            "version": int(time.time() * 1000) % 1000000,
            "grid_df": pd.DataFrame([
                {"商品コード": "", "商品名": "", "倉庫名": "", "受注残数(数量)": 0, "備考": ""}
                for _ in range(10)
            ])
        }

    if "slip_editor_state" not in st.session_state:
        st.session_state["slip_editor_state"] = create_blank_slip(next_order_no_str)

    slip_state = st.session_state["slip_editor_state"]
    v_key = slip_state.get("version", 1)

    @st.dialog("🔍 登録済み伝票を検索・呼び出し", width="large")
    def open_yayoi_search_modal():
        st.caption("💡 登録されている伝票を受注番号ごとに一覧表示しています。行を選択して「編集画面に開く」を押すと入力フォームに呼び出せます。")
        kw_modal = st.text_input("🔍 受注番号・得意先名・商品コードで検索", placeholder="例: 00002720 / 友和 / 丸藤")

        df_search_src = df_yayoi.copy()
        if kw_modal:
            df_search_src = df_search_src[
                df_search_src["受注番号"].astype(str).str.contains(kw_modal, case=False, na=False) |
                df_search_src["得意先名"].str.contains(kw_modal, case=False, na=False) |
                df_search_src["商品コード"].str.contains(kw_modal, case=False, na=False)
            ]

        if df_search_src.empty:
            st.info("該当する伝票がありません。")
            return

        modal_summary = df_search_src.groupby("受注番号", as_index=False).agg(
            受注日=("受注日", "first"),
            納入期日=("納入期日", "first"),
            得意先名=("得意先名", "first"),
            担当者名=("担当者名", "first"),
            明細数=("商品コード", "count"),
            合計数量=("数量", "sum"),
            備考=("備考", "first")
        ).sort_values(by=["受注日", "受注番号"], ascending=[False, False]).reset_index(drop=True)

        ev = st.dataframe(
            modal_summary[["受注日", "受注番号", "得意先名", "担当者名", "明細数", "合計数量", "納入期日", "備考"]],
            hide_index=True,
            use_container_width=True,
            height=240,
            selection_mode="single-row",
            on_select="rerun"
        )

        s_rows = ev.get("selection", {}).get("rows", [])
        target_idx = s_rows[0] if s_rows else 0
        target_ono = modal_summary.iloc[target_idx]["受注番号"]
        target_lines = df_yayoi[df_yayoi["受注番号"].astype(str) == str(target_ono)]

        st.markdown(
            f"""
            <div style="background:#eff6ff; border-left:5px solid #2563eb; padding:8px 12px; border-radius:4px; margin:6px 0;">
                <b>📦 選択中の伝票明細</b> ｜ 受注番号: <b>{target_ono}</b> ／ 得意先: <b>{modal_summary.iloc[target_idx]['得意先名']}</b>
            </div>
            """,
            unsafe_allow_html=True
        )
        st.dataframe(
            target_lines[["商品コード", "商品名/摘要", "倉庫名", "数量", "納入期日", "備考"]],
            hide_index=True,
            use_container_width=True,
            height=160
        )

        if st.button("📂 この伝票を入力フォームに開く（内容確認・修正・削除）", type="primary", use_container_width=True):
            first_r = target_lines.iloc[0]
            grid_rows = []
            for _, lr in target_lines.iterrows():
                c_code = clean_str(lr["商品コード"])
                if c_code and c_code not in item_code_list:
                    item_code_list.append(c_code)
                p_name = item_name_map.get(normalize_code(c_code), clean_str(lr.get("商品名/摘要", "")))
                u_wh = item_wh_map.get(normalize_code(c_code), clean_str(lr.get("倉庫名", "")))
                grid_rows.append({
                    "商品コード": c_code,
                    "商品名": p_name,
                    "倉庫名": u_wh,
                    "受注残数(数量)": int(lr["数量"]),
                    "備考": clean_str(lr.get("備考", ""))
                })
            while len(grid_rows) < 10:
                grid_rows.append({"商品コード": "", "商品名": "", "倉庫名": "", "受注残数(数量)": 0, "備考": ""})

            matched_comp = next((c for c in comp_select_list if c and c.startswith(clean_str(first_r["得意先名"]))), clean_str(first_r["得意先名"]))
            try:
                od_dt = pd.to_datetime(first_r["受注日"]).date()
            except Exception:
                od_dt = datetime.date.today()
            try:
                dd_dt = pd.to_datetime(first_r["納入期日"]).date()
            except Exception:
                dd_dt = datetime.date.today()

            st_name = clean_str(first_r.get("担当者名", ""))
            for clean_st in STAFF_OPTIONS_CLEAN:
                if clean_st and clean_st == st_name:
                    st_name = clean_st
                    break

            st.session_state["slip_editor_state"] = {
                "order_no": str(target_ono),
                "order_date": od_dt,
                "delivery_date": dd_dt,
                "comp_label": matched_comp,
                "staff": st_name,
                "summary": clean_str(first_r.get("備考", "")),
                "version": int(time.time() * 1000) % 1000000,
                "grid_df": pd.DataFrame(grid_rows[:100])
            }
            st.rerun()

    top_b1, top_b2, top_b3 = st.columns([1.8, 2.4, 2.0])
    if top_b1.button("📄 新規作成（白紙リセット）", use_container_width=True):
        st.session_state["slip_editor_state"] = create_blank_slip(next_order_no_str)
        st.rerun()

    if top_b2.button("🔍 登録済み伝票を検索・呼び出し", type="primary", use_container_width=True):
        open_yayoi_search_modal()

    if top_b3.button("➕ 入力行を10行追加", use_container_width=True):
        cur_df = slip_state["grid_df"]
        if len(cur_df) < 100:
            add_cnt = min(10, 100 - len(cur_df))
            extra_df = pd.DataFrame([
                {"商品コード": "", "商品名": "", "倉庫名": "", "受注残数(数量)": 0, "備考": ""}
                for _ in range(add_cnt)
            ])
            slip_state["grid_df"] = pd.concat([cur_df, extra_df], ignore_index=True)
            st.rerun()

    st.markdown('<div class="yayoi-slip-header">💻 受注伝票入力（商品コードを選ぶと「商品名(C列)」と「倉庫名(U列)」が自動連動）</div>', unsafe_allow_html=True)
    st.markdown('<div class="yayoi-slip-body">', unsafe_allow_html=True)

    h1, h2, h3, h4 = st.columns([1.4, 1.4, 1.4, 1.8])
    in_ono = h1.text_input("受注番号（自動採番）", value=slip_state["order_no"], key=f"ono_{v_key}")
    in_odate = h2.date_input("受注日", value=slip_state["order_date"], key=f"odate_{v_key}")
    in_ddate = h3.date_input("納入期日", value=slip_state["delivery_date"], key=f"ddate_{v_key}")

    staff_list = list(STAFF_OPTIONS_CLEAN)
    if slip_state["staff"] not in staff_list:
        staff_list.append(slip_state["staff"])
    in_staff = h4.selectbox(
        "担当者（新規は白紙）",
        options=staff_list,
        index=staff_list.index(slip_state["staff"]),
        key=f"staff_{v_key}"
    )

    hc1, hc2 = st.columns([2.6, 2.4])
    if slip_state["comp_label"] not in comp_select_list:
        comp_select_list.insert(1, slip_state["comp_label"])
    in_comp_label = hc1.selectbox(
        f"得意先（全 {len(master_comps):,} 社対応 ／ 新規は白紙）",
        options=comp_select_list,
        index=comp_select_list.index(slip_state["comp_label"]),
        key=f"comp_{v_key}"
    )
    in_summary = hc2.text_input(
        "摘要 / 伝票全体メモ（新規は白紙）",
        value=slip_state["summary"],
        placeholder="例: コープさっぽろ 12/3w",
        key=f"sum_{v_key}"
    )

    st.markdown(f"##### 📝 内訳明細（全 {len(master_items):,} 商品対応 ／ 商品コードを選ぶと隣の「商品名」「倉庫名」が自動表示されます）")

    raw_edited_grid = st.data_editor(
        slip_state["grid_df"],
        num_rows="dynamic",
        use_container_width=True,
        height=320,
        key=f"yayoi_grid_{v_key}",
        column_config={
            "商品コード": st.column_config.SelectboxColumn(
                f"商品コード (全 {len(master_items):,} 点)",
                options=item_code_list,
                width="medium"
            ),
            "商品名": st.column_config.TextColumn(
                "商品名 (自動)",
                width="large",
                disabled=True
            ),
            "倉庫名": st.column_config.TextColumn(
                "倉庫名 (自動)",
                width="small",
                disabled=True
            ),
            "受注残数(数量)": st.column_config.NumberColumn(
                "受注残数 (数量)",
                min_value=0,
                max_value=999999,
                step=1,
                default=0,
                width="small"
            ),
            "備考": st.column_config.TextColumn(
                "備考 / 個別摘要",
                width="medium"
            )
        }
    )

    need_sync = False
    synced_rows = []
    for _, r_g in raw_edited_grid.iterrows():
        c_code = clean_str(r_g.get("商品コード", ""))
        cur_name = clean_str(r_g.get("商品名", ""))
        cur_wh = clean_str(r_g.get("倉庫名", ""))
        qty_v = int(to_num(r_g.get("受注残数(数量)", 0)))
        note_v = clean_str(r_g.get("備考", ""))

        if c_code:
            expected_name = item_name_map.get(normalize_code(c_code), item_name_map.get(c_code, ""))
            expected_wh = item_wh_map.get(normalize_code(c_code), item_wh_map.get(c_code, ""))
        else:
            expected_name = ""
            expected_wh = ""

        if cur_name != expected_name or cur_wh != expected_wh:
            need_sync = True

        synced_rows.append({
            "商品コード": c_code,
            "商品名": expected_name,
            "倉庫名": expected_wh,
            "受注残数(数量)": qty_v,
            "備考": note_v
        })

    if need_sync:
        slip_state["grid_df"] = pd.DataFrame(synced_rows)
        slip_state["version"] = int(time.time() * 1000) % 1000000
        st.rerun()

    st.markdown("</div>", unsafe_allow_html=True)

    b_col1, b_col2 = st.columns([3.2, 1.0])
    if b_col1.button("💾 この伝票を登録・保存する（クリックで確定）", type="primary", use_container_width=True):
        comp_clean = in_comp_label.split(" （")[0].strip()
        staff_clean = in_staff.strip()

        if not comp_clean:
            st.error("⚠️ 「得意先」を選択してください。")
        elif not staff_clean:
            st.error("⚠️ 「担当者」を選択してください。")
        else:
            valid_items = []
            for idx_r, row_g in raw_edited_grid.iterrows():
                if idx_r >= 100:
                    break
                p_code = clean_str(row_g.get("商品コード", ""))
                qty_v = int(to_num(row_g.get("受注残数(数量)", 0)))
                if p_code and qty_v > 0:
                    p_name = item_name_map.get(normalize_code(p_code), item_name_map.get(p_code, ""))
                    u_wh = item_wh_map.get(normalize_code(p_code), item_wh_map.get(p_code, ""))
                    valid_items.append({
                        "受注番号": clean_str(in_ono),
                        "得意先名": comp_clean,
                        "商品コード": p_code,
                        "商品名/摘要": p_name,
                        "受注残数": qty_v,
                        "数量": qty_v,
                        "受注日": in_odate.strftime("%Y-%m-%d"),
                        "納入期日": in_ddate.strftime("%Y-%m-%d"),
                        "備考": clean_str(row_g.get("備考", "")) or clean_str(in_summary),
                        "倉庫名": u_wh,
                        "担当者名": staff_clean
                    })

            if not valid_items:
                st.error("⚠️ 商品コードが選択され、かつ「受注残数(数量)」が1以上の行がありません。")
            else:
                df_new_slip = pd.DataFrame(valid_items)
                if os.path.exists(MANUAL_YAYOI_CSV):
                    df_exist = pd.read_csv(MANUAL_YAYOI_CSV).fillna("")
                    df_exist = df_exist[df_exist["受注番号"].astype(str) != clean_str(in_ono)]
                    df_out = pd.concat([df_exist, df_new_slip], ignore_index=True)
                else:
                    df_out = df_new_slip

                df_out.to_csv(MANUAL_YAYOI_CSV, index=False, encoding="utf-8-sig")
                st.cache_data.clear()
                new_next_no = f"{max(max_seq + 1, int(clean_str(in_ono)) if clean_str(in_ono).isdigit() else max_seq + 1) + 1:08d}"
                st.session_state["slip_editor_state"] = create_blank_slip(new_next_no)
                st.success(f"✅ 受注番号『{in_ono}』（得意先: {comp_clean} ／ 明細 {len(valid_items)} 行）を保存しました！")
                st.rerun()

    if b_col2.button("🗑️ この伝票を削除する", use_container_width=True):
        if os.path.exists(MANUAL_YAYOI_CSV):
            df_exist = pd.read_csv(MANUAL_YAYOI_CSV).fillna("")
            df_exist = df_exist[df_exist["受注番号"].astype(str) != clean_str(in_ono)]
            df_exist.to_csv(MANUAL_YAYOI_CSV, index=False, encoding="utf-8-sig")
            st.cache_data.clear()
            st.session_state["slip_editor_state"] = create_blank_slip(next_order_no_str)
            st.success(f"🗑️ 受注番号『{in_ono}』を削除しました。")
            st.rerun()

    st.markdown("---")
    st.subheader("📋 登録済み伝票一覧（手入力保存分 ＆ 確保データ）")
    if os.path.exists(MANUAL_YAYOI_CSV):
        try:
            df_registered = pd.read_csv(MANUAL_YAYOI_CSV).fillna("")
            if not df_registered.empty:
                st.caption(f"💾 システム内で登録・保存された伝票（全 {df_registered['受注番号'].nunique()} 件）")
                st.dataframe(
                    df_registered[["受注日", "受注番号", "得意先名", "商品コード", "商品名/摘要", "数量", "倉庫名", "担当者名", "納入期日", "備考"]],
                    hide_index=True,
                    use_container_width=True,
                    height=260
                )
            else:
                st.info("手入力で新しく登録された伝票はまだありません。上のフォームから登録できます。")
        except Exception:
            st.info("登録済み伝票の読み込み待機中...")
    else:
        st.info("手入力で新しく登録された伝票はまだありません。上のフォームから登録できます。")

with tab4:
    st.subheader("🔍 売上CSV 品番別・ファイル別 集計明細チェック")
    if df_sales_detail.empty:
        st.info("現在、売上CSVが読み込まれていません。画面右上のエクスパンダーから yayoi_*.csv を選択してください。")
    else:
        chk_c1, chk_c2 = st.columns([1, 2])
        kw_sales_chk = chk_c1.text_input("品番で売上集計を検索（例: AHL-1400 / ASC-001）", value="ASC-001")
        df_s_filt = df_sales_detail[df_sales_detail["code_norm"].str.contains(kw_sales_chk.strip().upper(), na=False)] if kw_sales_chk else df_sales_detail
        chk_c2.metric(f"『{kw_sales_chk or '全商品'}』の集計売れ数 合計", f"{int(df_s_filt['qty'].sum()):,} 個", f"該当明細行数: {len(df_s_filt)} 行")
        st.dataframe(df_s_filt, hide_index=True, use_container_width=True, height=220)

    st.divider()
    ch1, ch2 = st.columns(2)
    with ch1:
        st.subheader("🚢 入荷予定あり・発注済みの商品一覧")
        inc = m_df[(m_df["入荷予定数_ロジザード"] > 0) | (m_df["発注日"] != "") | (m_df["入荷予定日_ロジザード"] != "")].copy()
        st.dataframe(inc[["品番", "商品区分", "在庫数_ロジザード", "入荷予定数_ロジザード", "入荷予定日_ロジザード", "発注日", "前回決定事項", "決定事項"]], hide_index=True, use_container_width=True, height=450)
    with ch2:
        st.subheader("📜 過去の会議シート 決定事項 全履歴")
        h_kw = st.text_input("品番・決定内容で検索", placeholder="例: 発注 / 様子見")
        hv = df_hist.sort_values(["会議日", "品番"], ascending=[False, True]) if not df_hist.empty else pd.DataFrame()
        if not hv.empty and h_kw:
            hv = hv[hv["品番"].str.contains(h_kw, case=False, na=False) | hv["決定事項"].str.contains(h_kw, case=False, na=False)]
        st.dataframe(hv, hide_index=True, use_container_width=True, height=450)
